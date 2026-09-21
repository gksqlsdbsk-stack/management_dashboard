"""엑셀/CSV 업로드 파싱·검증 [06]. DB·HTTP에 의존하지 않는 순수 함수 모음.

`parse_upload()`이 (항목, 프로젝트명, 값) 목록을 돌려주고, 오류는 예외로 알린다.
- UploadFileError: 파일 자체 오류 (확장자, 크기, 행 수, 인코딩, 손상 등)
- UploadRowsError: 행별 오류 목록 (한 건도 반영하지 않는다)
"""

import csv
import io
import re
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from .models import PROJECT_NAME_MAX_LENGTH

# 파일 형식·제한 [A-23]
ALLOWED_EXTENSIONS = (".xlsx", ".csv")
MAX_FILE_SIZE = 5 * 1024 * 1024
MAX_DATA_ROWS = 1000

ITEM_HEADER = "항목명"
PROJECT_HEADER = "프로젝트명"
VALUE_HEADER = "값"
REQUIRED_HEADERS = (ITEM_HEADER, PROJECT_HEADER, VALUE_HEADER)

_DECIMAL_PLACES = Decimal("0.0001")  # 소수 4자리 [A-03]
_MAX_DIGITS = 20
# 1500000, 1,500,000, -3, +3, 12.5 — 통화 기호·단위 문자·지수 표기는 불가
_NUMBER_PATTERN = re.compile(r"^[+-]?([0-9]+|[0-9]{1,3}(,[0-9]{3})+)(\.[0-9]+)?$")


class UploadFileError(Exception):
    """파일 자체 오류. `str(error)`가 사용자에게 보여줄 사유다."""


class UploadRowsError(Exception):
    def __init__(self, errors):
        super().__init__(f"{len(errors)}건의 오류")
        self.errors = errors  # [{"row": 3, "message": "3행: …"}]


# ---- 파일 읽기 ----

def validate_file(filename, size):
    lowered = (filename or "").lower()
    if not lowered.endswith(ALLOWED_EXTENSIONS):
        raise UploadFileError("엑셀(.xlsx) 또는 CSV(.csv) 파일만 업로드할 수 있습니다.")
    if size > MAX_FILE_SIZE:
        raise UploadFileError("파일 크기는 5MB 이하여야 합니다.")


def _iter_csv_rows(content):
    try:
        text = content.decode("utf-8-sig")  # BOM 허용 [A-23]
    except UnicodeDecodeError:
        raise UploadFileError("CSV 파일은 UTF-8 인코딩이어야 합니다.")
    try:
        yield from csv.reader(io.StringIO(text, newline=""))
    except csv.Error:
        raise UploadFileError("CSV 파일을 읽을 수 없습니다.")


def _iter_xlsx_rows(content):
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:  # 손상·비표준 파일은 종류가 다양해 한 가지 오류로 안내한다
        raise UploadFileError("엑셀 파일을 읽을 수 없습니다.")
    try:
        # 첫 번째 시트만 사용한다 [A-23]
        yield from workbook.worksheets[0].iter_rows(values_only=True)
    except Exception:
        raise UploadFileError("엑셀 파일을 읽을 수 없습니다.")
    finally:
        workbook.close()


def iter_rows(filename, content):
    """(행 번호, 셀 값 목록)을 차례로 돌려준다. 행 번호는 헤더가 1이다."""
    rows = _iter_csv_rows(content) if filename.lower().endswith(".csv") else _iter_xlsx_rows(content)
    for row_number, row in enumerate(rows, start=1):
        yield row_number, list(row)


# ---- 셀 해석 ----

def _text(cell):
    """항목명·프로젝트명 셀을 문자열로. 앞뒤 공백은 제거한다 [A-04, A-48]."""
    if cell is None:
        return ""
    if isinstance(cell, float) and cell.is_integer():
        return str(int(cell))  # 엑셀의 숫자 프로젝트명(2024.0) 대응
    return str(cell).strip()


def parse_decimal(cell):
    """값 셀을 Decimal로. 실패하면 ValueError(사용자 메시지)."""
    not_number = ValueError("값이 숫자가 아닙니다.")
    out_of_range = ValueError("값이 허용 범위를 벗어났습니다.")

    if isinstance(cell, bool):
        raise not_number
    if isinstance(cell, int):
        value = Decimal(cell)
    elif isinstance(cell, float):
        try:
            value = Decimal(repr(cell))
        except InvalidOperation:
            raise not_number
        if not value.is_finite():
            raise not_number
    elif isinstance(cell, Decimal):
        value = cell
    elif isinstance(cell, str):
        stripped = cell.strip()
        if not _NUMBER_PATTERN.match(stripped):
            raise not_number
        value = Decimal(stripped.replace(",", ""))
    else:
        raise not_number

    try:
        quantized = value.quantize(_DECIMAL_PLACES)
    except InvalidOperation:  # 자릿수가 너무 크다
        raise out_of_range
    if quantized != value or len(quantized.as_tuple().digits) > _MAX_DIGITS:
        raise out_of_range
    return quantized


# ---- 검증 ----

def _topic_particle(word):
    """주제 조사 은/는. 한글 끝 글자의 받침으로 고르고, 그 외 문자로 끝나면 '은(는)'."""
    last = word[-1]
    if "가" <= last <= "힣":
        return "은" if (ord(last) - ord("가")) % 28 else "는"
    return "은(는)"


def _header_index(header_row):
    index = {}
    for position, cell in enumerate(header_row):
        index.setdefault(_text(cell), position)
    return index


def _cell(row, position):
    return row[position] if position < len(row) else None


def parse_upload(filename, content, items):
    """파일을 읽어 검증하고 [(item, project_name, Decimal)] 을 돌려준다 [06 §4].

    `items`: 대상 부서의 활성 입력 항목(InputItem) 목록.
    """
    validate_file(filename, len(content))
    rows = iter_rows(filename, content)

    header = next(rows, None)
    header_index = _header_index(header[1]) if header else {}
    header_errors = [
        {"row": 1, "message": f"헤더에 '{name}' 열이 없습니다."} for name in REQUIRED_HEADERS if name not in header_index
    ]
    if header_errors:
        raise UploadRowsError(header_errors)
    item_col, project_col, value_col = (header_index[name] for name in REQUIRED_HEADERS)

    items_by_name = {item.name: item for item in items}
    errors = []
    parsed = []
    first_seen = {}  # (item id, 프로젝트명) → 처음 나온 행 번호
    data_rows = 0

    for row_number, row in rows:
        raw_name = _cell(row, item_col)
        raw_project = _cell(row, project_col)
        raw_value = _cell(row, value_col)
        name, project_name = _text(raw_name), _text(raw_project)
        value_blank = raw_value is None or (isinstance(raw_value, str) and not raw_value.strip())
        if not name and not project_name and value_blank:
            continue  # 세 열이 모두 빈 행은 무시 [A-48]

        data_rows += 1
        if data_rows > MAX_DATA_ROWS:
            raise UploadFileError(f"데이터는 최대 {MAX_DATA_ROWS:,}행까지 업로드할 수 있습니다.")

        row_errors = []

        item = items_by_name.get(name)
        if not name:
            row_errors.append(f"{row_number}행: 항목명이 비어 있습니다.")
        elif item is None:
            row_errors.append(f"{row_number}행: '{name}' 항목을 찾을 수 없습니다.")
        elif item.scope == "MONTHLY" and project_name:
            row_errors.append(f"{row_number}행: '{name}'{_topic_particle(name)} 프로젝트명을 입력하지 않는 항목입니다.")
        elif item.scope == "PROJECT" and not project_name:
            row_errors.append(f"{row_number}행: '{name}'{_topic_particle(name)} 프로젝트명이 필요합니다.")
        if len(project_name) > PROJECT_NAME_MAX_LENGTH:
            row_errors.append(f"{row_number}행: 프로젝트명이 너무 깁니다.")

        value = None
        try:
            value = parse_decimal(raw_value)
        except ValueError as error:
            row_errors.append(f"{row_number}행: {error}")

        if not row_errors:
            key = (item.id, project_name)
            if key in first_seen:
                row_errors.append(f"{row_number}행: {first_seen[key]}행과 중복된 항목입니다.")
            else:
                first_seen[key] = row_number
                parsed.append((item, project_name, value))

        errors.extend({"row": row_number, "message": message} for message in row_errors)

    if errors:
        raise UploadRowsError(errors)
    if not parsed:
        raise UploadFileError("업로드할 데이터 행이 없습니다.")
    return parsed
